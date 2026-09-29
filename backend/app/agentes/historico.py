"""
A linha do tempo do caso — tudo que foi feito, na ordem em que foi feito.

O SINTOMA
---------
A perícia foi realizada, o responsável anotou o resultado na agenda, e no
card do cliente não apareceu nada. O mesmo com a pendência resolvida, a
tarefa concluída, o prazo cumprido, a fase que avançou.

A CAUSA
-------
Cada ação grava um evento em `eventos` — há mais de noventa tipos
espalhados pelos agentes. Mas o card mostrava apenas a tabela
`mensagens`, e o único lugar que lia `eventos` era o e-mail de prestação
de contas, através de um dicionário com dezesseis rótulos: tudo que não
estivesse nessa lista era descartado em silêncio.

Ou seja: o sistema registrava o trabalho e não tinha onde mostrá-lo.

O QUE ESTE MÓDULO FAZ
---------------------
Traduz os eventos para português de quem lê, e os separa em três
naturezas, porque misturá-las é o que tornava o histórico ilegível:

  TRABALHO     o que a equipe fez para resolver a demanda. É disto que
               o relatório do caso é feito.
  CLIENTE      conversa, envio de documento, ciência. Interação, não
               produção.
  SISTEMA      rotina automática e erro técnico. Fica guardado para
               auditoria e não entra no relatório do cliente.

O relatório é a memória do caso. Vale para prestar contas, para provar
diligência num pedido de honorários, e para quem pegar o caso daqui a
dois anos entender o que já foi tentado.
"""
from __future__ import annotations

from datetime import datetime

from ..core.db import get_db

TRABALHO, CLIENTE, SISTEMA = "TRABALHO", "CLIENTE", "SISTEMA"

# Cada evento: (natureza, rótulo legível).
#
# A lista é longa de propósito. A anterior tinha dezesseis entradas e
# descartava o resto — e era justamente no resto que estava o trabalho:
# audiência realizada, pendência resolvida, peça redigida, prazo
# cumprido. Um histórico que esconde o trabalho não serve para prestar
# contas de coisa nenhuma.
EVENTOS: dict[str, tuple[str, str]] = {
    # ── Entrada do caso ──────────────────────────────────────────
    "TRIAGEM": (TRABALHO, "Triagem do caso"),
    "CASO_ESCRITORIO_CRIADO": (TRABALHO, "Caso cadastrado pelo escritório"),
    "PROCESSO_INICIADO_ESTEIRA": (TRABALHO, "Caso entrou na esteira"),
    "PROCESSO_IMPORTADO": (TRABALHO, "Processo importado dos tribunais"),
    "ESTADO_MUDOU": (TRABALHO, "Mudança de fase"),
    "CASO_SUSPENSO": (TRABALHO, "Caso suspenso"),
    "CASO_ARQUIVADO": (TRABALHO, "Caso arquivado"),
    "CASO_ATIVADO": (TRABALHO, "Caso reativado"),
    "CASO_PARADO_SINALIZADO": (SISTEMA, "Caso sinalizado como parado"),
    "CASO_ARQUIVADO_LEAD_FRIO": (TRABALHO, "Arquivado por falta de retorno"),
    "CASO_RETOMADO": (TRABALHO, "Caso retomado"),

    # ── Honorários e contrato de serviço ─────────────────────────
    "CONTRATO_SERVICO_INICIADO": (TRABALHO, "Proposta de honorários iniciada"),
    "HONORARIOS_ALTERADOS": (TRABALHO, "Honorários ajustados"),
    "HONORARIOS_LIDOS_DA_CONVERSA": (SISTEMA, "Honorários lidos da conversa"),
    "DADOS_CONTRATO_REGISTRADOS": (TRABALHO, "Dados do contrato registrados"),
    "CONTRATO_ENVIADO": (TRABALHO, "Contrato enviado ao cliente"),
    "COBRANCA_CRIADA": (TRABALHO, "Cobrança emitida"),

    # ── Documentos ───────────────────────────────────────────────
    "DOCUMENTO_GERADO": (TRABALHO, "Documento elaborado"),
    "DOCUMENTO_EDITADO": (TRABALHO, "Documento editado"),
    "DOCUMENTO_AJUSTADO": (TRABALHO, "Documento ajustado"),
    "DOCUMENTO_REVISADO": (TRABALHO, "Documento revisado"),
    "DOCUMENTO_ATUALIZADO_INTEGRAL": (TRABALHO, "Documento refeito a partir do modelo"),
    "DOCUMENTO_CONVERTIDO_PDF": (SISTEMA, "Documento convertido em PDF"),
    "DOCUMENTO_ANEXADO": (TRABALHO, "Documento anexado ao caso"),
    "DOCUMENTOS_ENVIADOS_LOTE": (TRABALHO, "Documentos anexados"),
    "DOCUMENTO_APROVADO": (TRABALHO, "Documento aprovado"),
    "DOCUMENTO_ENVIADO_ASSINATURA": (TRABALHO, "Enviado para assinatura"),
    "DOCUMENTO_ENVIADO_CLIENTE": (CLIENTE, "Documento enviado ao cliente"),
    "DOCUMENTO_ASSINADO": (CLIENTE, "Documento assinado"),
    "DOCUMENTO_ASSINADO_RECEBIDO": (CLIENTE, "Documento assinado recebido"),
    "ASSINADO_RECEBIDO_EMAIL": (CLIENTE, "Assinado recebido por e-mail"),
    "DOCUMENTOS_RECEBIDOS_DO_CLIENTE": (CLIENTE, "Documentos recebidos do cliente"),
    "SOLICITACAO_CLIENTE": (TRABALHO, "Documento solicitado ao cliente"),
    "AGUARDANDO_CLIENTE": (TRABALHO, "Aguardando documento do cliente"),

    # ── Peça e protocolo ─────────────────────────────────────────
    "MINUTA_REDIGIDA": (TRABALHO, "Peça redigida"),
    "PETICAO_CRIADA": (TRABALHO, "Petição criada"),
    "PRECEDENTES_INJETADOS": (TRABALHO, "Jurisprudência pesquisada e inserida"),
    "PETICIONAMENTO_VALIDADO": (TRABALHO, "Peça validada antes do protocolo"),
    "PETICIONAMENTO_FORCADO": (TRABALHO, "Protocolo liberado por decisão do advogado"),

    # ── Andamento processual ─────────────────────────────────────
    "FASE_JUDICIAL": (TRABALHO, "Fase processual atualizada"),
    "FASE_JUDICIAL_MANUAL": (TRABALHO, "Fase corrigida à mão"),
    "FASE_RECEBIMENTO": (TRABALHO, "Movido para execução"),
    "PRAZO_CADASTRADO": (TRABALHO, "Prazo cadastrado"),
    "PRAZO_DISTRIBUIDO": (TRABALHO, "Prazo atribuído a um responsável"),
    "PRAZO_AJUSTADO": (TRABALHO, "Prazo corrigido à mão"),
    "PRAZO_CUMPRIDO": (TRABALHO, "Prazo cumprido"),
    "PRAZO_VENCIDO_EM_ABERTO": (SISTEMA, "Prazo venceu sem baixa"),
    "PRAZO_NAO_CRIADO": (SISTEMA, "Prazo não pôde ser criado"),
    "MONITORAMENTO_CRIADO": (SISTEMA, "Monitoramento do processo ligado"),

    # ── Agenda ───────────────────────────────────────────────────
    "AGENDA_CRIADA": (TRABALHO, "Compromisso marcado"),
    "AGENDA_REAGENDADA": (TRABALHO, "Compromisso remarcado"),
    "AGENDA_REALIZADA": (TRABALHO, "Compromisso realizado"),
    "AGENDA_CANCELADA": (TRABALHO, "Compromisso cancelado"),
    "AGENDA_RESPOSTA": (CLIENTE, "Resposta ao convite"),
    "COMPROMISSO_AGENDADO": (TRABALHO, "Compromisso agendado pela controladoria"),

    # ── Pendências e tarefas ─────────────────────────────────────
    "PENDENCIA_ANOTADA": (TRABALHO, "Pendência anotada"),
    "PENDENCIA_RESOLVIDA": (TRABALHO, "Pendência resolvida"),
    "PENDENCIA_REAGENDADA": (TRABALHO, "Pendência adiada"),
    "PENDENCIA_CANCELADA": (TRABALHO, "Pendência cancelada"),
    "TAREFA_CONCLUIDA": (TRABALHO, "Tarefa concluída"),
    "TAREFA_REAGENDADA": (TRABALHO, "Tarefa adiada"),
    "TAREFA_ATRIBUIDA": (TRABALHO, "Tarefa atribuída"),

    # ── Atendimento ──────────────────────────────────────────────
    "ATENDIMENTO_ABERTO": (TRABALHO, "Atendimento aberto"),
    "ATENDIMENTO_CLIENTE_ENTROU": (CLIENTE, "Cliente entrou no atendimento"),
    "ATENDIMENTO_ENCERRADO": (TRABALHO, "Atendimento encerrado"),
    "ATENDIMENTO_TRANSCRITO": (TRABALHO, "Atendimento transcrito"),
    "GRAVACAO_AUTORIZADA_DURANTE_ATENDIMENTO": (CLIENTE, "Cliente autorizou a gravação"),
    "GRAVACAO_DESCARTADA_SEM_CONSENTIMENTO": (SISTEMA, "Gravação descartada por falta de autorização"),
    "AUDIO_ATENDIMENTO_GUARDADO": (SISTEMA, "Áudio do atendimento guardado"),

    # ── Comunicação ──────────────────────────────────────────────
    "AVISO_ENVIADO": (CLIENTE, "Comunicação enviada"),
    "CIENCIA_CLIENTE": (CLIENTE, "Cliente deu ciência"),
    "CLIENTE_RETORNOU": (CLIENTE, "Cliente respondeu"),
    "RESPOSTA_CLIENTE_EMAIL": (CLIENTE, "Cliente respondeu por e-mail"),
    "CONTATO_ATUALIZADO": (SISTEMA, "Contato atualizado"),
    "PRESTACAO_DE_CONTAS": (TRABALHO, "Prestação de contas enviada"),
    "ANIVERSARIOS_ENVIADOS": (SISTEMA, "Mensagem de aniversário"),

    # ── Ruído técnico ────────────────────────────────────────────
    "ERRO_AGENTE": (SISTEMA, "Falha de um agente"),
    "ERRO_FERRAMENTA": (SISTEMA, "Falha de ferramenta"),
    "WEBHOOK_ZAPSIGN": (SISTEMA, "Retorno da assinatura eletrônica"),
    "WEBHOOK_ASAAS": (SISTEMA, "Retorno do meio de pagamento"),
    "WEBHOOK_ESCAVADOR": (SISTEMA, "Retorno do monitoramento"),
    "CONTROLADORIA_RODOU": (SISTEMA, "Varredura da controladoria"),
    "PLANO_DIARIO": (SISTEMA, "Plano do dia montado"),
    "PLANO_SEMANAL": (SISTEMA, "Plano da semana montado"),
    "PLANO_REAJUSTADO": (SISTEMA, "Plano reajustado"),
    "RADAR_DATAJUD": (SISTEMA, "Radar de jurisprudência"),
    "CEREBRO_CLOUD": (SISTEMA, "Consulta ao banco de teses"),
    "TESE_EXCLUIDA_OVERRULING": (SISTEMA, "Tese retirada por superação"),
}


def _quando(e: dict) -> str:
    return (e.get("criado_em") or "")[:19].replace("T", " ")


def _detalhe(tipo: str, p: dict) -> str:
    """A frase que explica o evento, montada do payload.

    Sem isto, "Compromisso realizado" não diz o que aconteceu na
    audiência — e é exatamente o que interessa saber depois."""
    if not isinstance(p, dict):
        return ""
    for chave in ("resultado", "o_que_foi_feito", "nota", "motivo",
                  "descricao", "titulo", "observacao", "texto", "nome"):
        v = p.get(chave)
        if isinstance(v, str) and v.strip():
            base = v.strip()
            break
    else:
        base = ""

    extras = []
    if tipo == "ESTADO_MUDOU" and p.get("de"):
        extras.append(f"de {p['de']} para {p.get('para', '')}")
    if tipo in ("AGENDA_REAGENDADA", "PENDENCIA_REAGENDADA", "TAREFA_REAGENDADA"):
        if p.get("de") and p.get("para"):
            extras.append(f"de {p['de']} para {p['para']}")
    if tipo == "PRAZO_AJUSTADO" and p.get("de"):
        extras.append(f"de {p['de']} para {p.get('para', '')}")
    if p.get("tipo") and tipo.startswith("AGENDA"):
        extras.append(str(p["tipo"]).lower())
    if p.get("data"):
        extras.append(f"data {p['data']}")

    partes = [x for x in [base, "; ".join(extras)] if x]
    return " — ".join(partes)[:400]


def _quem(p: dict) -> str:
    if not isinstance(p, dict):
        return ""
    for chave in ("quem", "por", "autor", "responsavel", "usuario"):
        v = p.get(chave)
        if isinstance(v, str) and v.strip():
            return v.strip()[:80]
    return ""


def linha_do_tempo(caso_id: str, natureza: str | None = None,
                   incluir_conversas: bool = True) -> list[dict]:
    """Tudo que aconteceu no caso, do mais antigo para o mais recente.

    `natureza` filtra: TRABALHO devolve só o que a equipe fez, que é o
    que o relatório usa.
    """
    db = get_db()
    itens: list[dict] = []

    eventos = db.table("eventos").select("tipo,payload,criado_em") \
        .eq("caso_id", caso_id).order("criado_em").limit(1000).execute().data or []
    for e in eventos:
        tipo = e.get("tipo") or ""
        nat, rotulo = EVENTOS.get(tipo, (SISTEMA, tipo.replace("_", " ").capitalize()))
        p = e.get("payload") or {}
        itens.append({
            "em": _quando(e), "natureza": nat, "titulo": rotulo,
            "detalhe": _detalhe(tipo, p), "quem": _quem(p),
            "origem": "evento", "tipo": tipo,
        })

    if incluir_conversas:
        msgs = db.table("mensagens").select("autor,canal,conteudo,criado_em") \
            .eq("caso_id", caso_id).order("criado_em").limit(1000).execute().data or []
        for m in msgs:
            autor = (m.get("autor") or "").upper()
            # Nota interna é trabalho: é o advogado registrando o que
            # fez ou decidiu. Mensagem do cliente é interação.
            interna = autor in ("HUMANO", "CRM", "ESCRITORIO")
            itens.append({
                "em": _quando(m),
                "natureza": TRABALHO if interna else CLIENTE,
                "titulo": "Anotação interna" if interna else "Mensagem do cliente",
                "detalhe": (m.get("conteudo") or "")[:600],
                "quem": autor.title(), "origem": "mensagem",
                "tipo": "NOTA" if interna else "MENSAGEM",
            })

    itens.sort(key=lambda x: x["em"])
    if natureza:
        itens = [i for i in itens if i["natureza"] == natureza.upper()]
    return itens


def resumo(caso_id: str) -> dict:
    """Quanto trabalho existe registrado, e desde quando."""
    tudo = linha_do_tempo(caso_id)
    trabalho = [i for i in tudo if i["natureza"] == TRABALHO]
    return {
        "total": len(tudo),
        "trabalho": len(trabalho),
        "cliente": len([i for i in tudo if i["natureza"] == CLIENTE]),
        "sistema": len([i for i in tudo if i["natureza"] == SISTEMA]),
        "primeiro": tudo[0]["em"] if tudo else None,
        "ultimo": tudo[-1]["em"] if tudo else None,
    }


# ── Relatório do caso ───────────────────────────────────────────
def relatorio_docx(caso_id: str, so_trabalho: bool = True) -> tuple[bytes, str]:
    """O relatório do caso em Word, pronto para entregar.

    Por padrão traz só o TRABALHO: o cliente quer saber o que foi feito
    para resolver a demanda dele, não quantas vezes a controladoria
    varreu o banco de madrugada.
    """
    import io
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    db = get_db()
    caso = db.table("casos").select("*,clientes(nome,cpf_cnpj,email)") \
        .eq("id", caso_id).limit(1).execute().data
    if not caso:
        raise ValueError("Caso não encontrado.")
    caso = caso[0]
    cli = caso.get("clientes") or {}

    itens = linha_do_tempo(caso_id, natureza=TRABALHO if so_trabalho else None)

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("RELATÓRIO DE ATIVIDADES DO CASO")
    r.bold = True
    r.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("FSC Advocacia — Fábio Silva Cunha (OAB/RO 10.849)").italic = True

    doc.add_paragraph()
    ident = doc.add_table(rows=0, cols=2)
    ident.style = "Table Grid"
    def linha(a, b):
        c = ident.add_row().cells
        c[0].text = a
        c[1].text = str(b or "—")
        for p in c[0].paragraphs:
            for run in p.runs:
                run.bold = True
    linha("Cliente", cli.get("nome"))
    linha("CPF/CNPJ", cli.get("cpf_cnpj"))
    linha("Atendimento", caso.get("numero_atendimento"))
    linha("Processo", caso.get("numero_processo"))
    linha("Matéria", caso.get("grupo"))
    linha("Fase atual", caso.get("estado"))
    linha("Aberto em", (caso.get("criado_em") or "")[:10])
    linha("Relatório emitido em", datetime.now().strftime("%d/%m/%Y %H:%M"))

    doc.add_paragraph()
    h = doc.add_paragraph()
    h.add_run("O QUE FOI REALIZADO").bold = True

    if not itens:
        doc.add_paragraph(
            "Não há atividades registradas para este caso até a presente data.")
    else:
        doc.add_paragraph(
            f"Estão registradas {len(itens)} atividades, da mais antiga para a "
            f"mais recente.").italic = True
        tab = doc.add_table(rows=1, cols=3)
        tab.style = "Table Grid"
        cab = tab.rows[0].cells
        for i, titulo in enumerate(["Data", "Atividade", "Detalhe"]):
            cab[i].text = titulo
            for p in cab[i].paragraphs:
                for run in p.runs:
                    run.bold = True
        for it in itens:
            c = tab.add_row().cells
            c[0].text = it["em"][:16]
            c[1].text = it["titulo"]
            det = it["detalhe"] or ""
            if it["quem"]:
                det = (det + f" (por {it['quem']})").strip()
            c[2].text = det

    doc.add_paragraph()
    rodape = doc.add_paragraph()
    rodape.add_run(
        "Este relatório é gerado a partir do registro automático de atividades "
        "do sistema de gestão do escritório. Cada linha corresponde a uma ação "
        "efetivamente executada e registrada na data indicada."
    ).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    nome = f"relatorio-{caso.get('numero_atendimento') or caso_id[:8]}.docx"
    return buf.getvalue(), nome
