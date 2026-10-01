"""
O PAPEL TIMBRADO DO ESCRITÓRIO.

As duas imagens saíram do arquivo que a equipe já usa à mão, o
PAPEL TIMBRADO-FABIO.docx: a faixa com o brasão no alto e a tarja de
contatos no pé. Ficam versionadas junto com o código, e não numa pasta
compartilhada, porque timbre em pasta compartilhada é timbre que um dia
alguém troca por outra versão e ninguém consegue dizer qual saiu em
cada contrato assinado.

NO CABEÇALHO DA SEÇÃO, NÃO NO CORPO

É isso que faz o timbre repetir em todas as páginas, que é o que
distingue papel timbrado de folha com um logo colado em cima. Posto no
corpo, um contrato de dezesseis folhas sairia timbrado em uma, e as
outras quinze chegariam ao cliente parecendo rascunho de impressora.

A FAIXA SANGRA PARA FORA DA MARGEM

Ela é mais larga que a área útil, de propósito, e o recuo negativo a
empurra até a borda da folha. Sem isso ela ficaria encolhida no meio da
página, com duas tiras brancas nas laterais, que é exatamente o que
denuncia timbre improvisado.

ALTURA, E O ERRO QUE ELA ESCONDE

A faixa tem 3,3 cm na largura da folha. A margem superior precisa caber
ela inteira mais um respiro: com margem curta, o primeiro parágrafo
nasce por cima do brasão. Esse defeito não aparece no editor, só no PDF
convertido, que é quando o documento já está a um clique do cliente.
"""
from __future__ import annotations

from pathlib import Path

PASTA = Path(__file__).resolve().parent.parent / "modelos" / "marca"
TOPO = PASTA / "cabecalho.png"
PE = PASTA / "rodape.png"

# Centímetros. A largura do topo é a da folha A4 inteira.
TOPO_CM = 21.0
PE_CM = 6.5

# As medidas de página de um documento timbrado. Ficam aqui para que
# quem acrescentar um gerador novo não precise descobrir de novo que a
# margem superior tem de ser maior que a faixa.
MARGEM_TOPO_CM = 4.2
# A margem de pé cabe a tarja de contatos MAIS a linha "fl. X de Y". Com
# 2,4 cm, que era a medida de antes da numeração, o último bloco de
# assinatura da folha era impresso por cima da tarja: o nome do cliente
# atravessado pelo telefone do escritório, e isso só aparece no PDF.
MARGEM_PE_CM = 3.1
DISTANCIA_CABECALHO_CM = 0.4
DISTANCIA_RODAPE_CM = 0.8


def disponivel() -> bool:
    return TOPO.exists() and PE.exists()


def aplicar(sec) -> None:
    """Põe a faixa no cabeçalho e a tarja no rodapé de uma seção.

    Falhar aqui não derruba o documento: um contrato sem timbre ainda é
    um contrato, e um contrato que não sai não é nada, com o cliente
    esperando do outro lado."""
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    sec.header_distance = Cm(DISTANCIA_CABECALHO_CM)
    sec.top_margin = Cm(MARGEM_TOPO_CM)
    sec.footer_distance = Cm(DISTANCIA_RODAPE_CM)
    sec.bottom_margin = Cm(MARGEM_PE_CM)

    def _por(parte, caminho, largura_cm, recuo_cm, alinhamento):
        par = parte.paragraphs[0] if parte.paragraphs else parte.add_paragraph()
        par.alignment = alinhamento
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.left_indent = Cm(recuo_cm)
        par.paragraph_format.right_indent = Cm(recuo_cm)
        par.add_run().add_picture(str(caminho), width=Cm(largura_cm))

    try:
        # O recuo negativo é a própria margem esquerda, em espelho: a
        # imagem começa onde a folha começa.
        _por(sec.header, TOPO, TOPO_CM, -float(sec.left_margin.cm),
             WD_ALIGN_PARAGRAPH.LEFT)
    except Exception as e:
        print(f"[timbre] faixa do cabeçalho não entrou: {e}")

    try:
        _por(sec.footer, PE, PE_CM, 0, WD_ALIGN_PARAGRAPH.CENTER)
    except Exception as e:
        print(f"[timbre] tarja do rodapé não entrou: {e}")
